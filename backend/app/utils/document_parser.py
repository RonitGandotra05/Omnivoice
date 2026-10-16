"""
Document Parser – parses PDF, TXT, MD, JSON, DOCX into plain text.
"""

import json
from pathlib import Path
from io import BytesIO

# Optional imports (graceful degradation)
try:
    from pypdf import PdfReader
    _HAS_PDF = True
except ImportError:
    _HAS_PDF = False

try:
    import docx
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False


SUPPORTED_EXTENSIONS = {"txt", "md", "json"}
if _HAS_PDF:
    SUPPORTED_EXTENSIONS.add("pdf")
if _HAS_DOCX:
    SUPPORTED_EXTENSIONS.add("docx")


def get_supported_types() -> set[str]:
    return SUPPORTED_EXTENSIONS


# ── Public API ────────────────────────────────────────────────────────

def parse_buffer(buffer: bytes, content_type: str, filename: str = "") -> dict:
    """
    Parse a file buffer into plain text.

    Returns:
        {"content": str, "metadata": dict}
    """
    ext = _normalize_type(content_type, filename)
    from app.security.uploads import validate_passive
    validate_passive(buffer, ext)

    parsers = {
        "pdf": _parse_pdf,
        "txt": _parse_text,
        "md": _parse_text,
        "json": _parse_json,
        "docx": _parse_docx,
    }

    parser = parsers.get(ext)
    if parser is None:
        raise ValueError(f"Unsupported file type: {ext}")

    result = parser(buffer, filename)
    from app.config.settings import get_settings
    if len(result['content']) > get_settings().max_document_characters:
        raise ValueError('Extracted text exceeds the document limit. Split this document.')
    return result


# ── Parsers ───────────────────────────────────────────────────────────

def _parse_pdf(buffer: bytes, filename: str) -> dict:
    if not _HAS_PDF:
        raise ValueError("PDF support is not installed.")
    from app.config.settings import get_settings
    from app.utils.ocr import recognize, availability
    settings = get_settings()
    reader = PdfReader(BytesIO(buffer))
    if reader.is_encrypted:
        raise ValueError("Password-protected PDFs are not supported. Upload an unlocked copy.")
    if not 1 <= len(reader.pages) <= settings.pdf_max_pages:
        raise ValueError(f"PDFs must have 1–{settings.pdf_max_pages} pages. Split larger documents.")
    from app.security.uploads import validate_pdf_objects
    validate_pdf_objects(reader)
    pages, quality, warnings = [], [], []
    rendered = None
    ocr_count = 0
    try:
        for index, page in enumerate(reader.pages):
            text = _clean(page.extract_text() or "")
            readable = len(text) >= 40 and sum(c.isalnum() for c in text) / max(1, len(text)) >= 0.4 and "\ufffd" not in text
            info = {"page": index + 1, "method": "text", "confidence": None}
            if not readable:
                if not settings.ocr_enabled or not availability(settings.ocr_languages):
                    raise ValueError(f"Page {index + 1} requires OCR. Install Tesseract and the configured language packs, or upload a searchable PDF.")
                ocr_count += 1
                if ocr_count > settings.pdf_max_ocr_pages:
                    raise ValueError(f"At most {settings.pdf_max_ocr_pages} scanned pages per file are supported. Split this PDF.")
                if rendered is None:
                    import pypdfium2 as pdfium
                    rendered = pdfium.PdfDocument(buffer)
                rendered_page = rendered[index]
                scale = settings.ocr_dpi / 72
                width, height = rendered_page.get_size()
                if width * height * scale * scale > settings.pdf_max_render_pixels:
                    rendered_page.close()
                    raise ValueError("A PDF page exceeds the safe image size. Resize or split the document.")
                bitmap = rendered_page.render(scale=scale)
                image = bitmap.to_pil().convert('RGB')
                try:
                    result = recognize(image, languages=settings.ocr_languages)
                finally:
                    image.close()
                    bitmap.close()
                    rendered_page.close()
                text = _clean(result['text'])
                info.update(method="ocr", confidence=result['confidence'], rotation=result['rotation'])
                if len(text) < 10:
                    warnings.append(f"Page {index + 1}: no reliable text could be extracted.")
                elif result['confidence'] < settings.ocr_min_confidence:
                    warnings.append(f"Page {index + 1}: OCR confidence is low; review the extracted text.")
            pages.append(text)
            quality.append(info)
            if sum(map(len, pages)) > settings.max_document_characters:
                raise ValueError("Extracted text exceeds the document limit. Split this PDF.")
    finally:
        if rendered is not None:
            rendered.close()
    content = _clean("\n\n".join(pages))
    if len(content) < 10:
        raise ValueError("No readable text was found. Provide a clearer scan or a text document.")
    return {"content": content, "metadata": {"pages": len(pages), "filename": filename,
        "ocr_pages": ocr_count, "page_quality": quality, "warnings": warnings,
        "needs_review": bool(warnings) or (ocr_count > 0 and settings.ocr_review_required)}}


def _parse_text(buffer: bytes, filename: str) -> dict:
    return {
        "content": _clean(buffer.decode("utf-8-sig")),
        "metadata": {"filename": filename},
    }


def _parse_json(buffer: bytes, filename: str) -> dict:
    data = json.loads(buffer.decode("utf-8-sig"))
    text = _json_to_text(data)
    return {
        "content": _clean(text),
        "metadata": {"filename": filename, "format": "json"},
    }


def _parse_docx(buffer: bytes, filename: str) -> dict:
    if not _HAS_DOCX:
        raise RuntimeError("python-docx not installed – cannot parse DOCX")
    import zipfile
    with zipfile.ZipFile(BytesIO(buffer)) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
            raise ValueError('DOCX expanded contents exceed the safe size limit.')
    doc = docx.Document(BytesIO(buffer))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    return {
        "content": _clean("\n\n".join(paragraphs)),
        "metadata": {"filename": filename, "paragraphs": len(paragraphs)},
    }


# ── Helpers ───────────────────────────────────────────────────────────

def _json_to_text(data) -> str:
    """Convert various JSON shapes into readable text."""
    if isinstance(data, list):
        parts = []
        for item in data:
            if isinstance(item, dict):
                if "question" in item and "answer" in item:
                    parts.append(f"Q: {item['question']}\nA: {item['answer']}")
                elif "title" in item and "content" in item:
                    parts.append(f"{item['title']}\n{item['content']}")
                else:
                    parts.append(json.dumps(item, indent=2))
            elif isinstance(item, str):
                parts.append(item)
        return "\n\n".join(parts)
    elif isinstance(data, dict):
        return "\n\n".join(
            f"{k}: {v}" if isinstance(v, str) else f"{k}:\n{json.dumps(v, indent=2)}"
            for k, v in data.items()
        )
    return str(data)


def _normalize_type(content_type: str, filename: str) -> str:
    mime_map = {
        "application/pdf": "pdf",
        "text/plain": "txt",
        "text/markdown": "md",
        "application/json": "json",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    }
    if content_type in mime_map:
        return mime_map[content_type]
    # Fall back to extension
    ext = Path(filename).suffix.lstrip(".").lower()
    return ext if ext else content_type


def _clean(text: str) -> str:
    """Normalise whitespace."""
    lines = text.replace("\r\n", "\n").replace("\t", " ").split("\n")
    lines = [ln.strip() for ln in lines]
    text = "\n".join(lines)
    # Collapse 3+ newlines
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    return text.strip()
